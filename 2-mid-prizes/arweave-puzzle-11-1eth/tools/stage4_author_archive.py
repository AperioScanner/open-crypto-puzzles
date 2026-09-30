#!/usr/bin/env python3
"""Mine the public @ArweaveP tweet archive mirrored by HomelessPhD for Puzzle #11 clues."""
from __future__ import annotations
import json, re
from datetime import datetime, date
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"analysis"/"runs"/"stage4-author-archive"
XLSX=Path("/tmp/ArweaveP_user_tweets.xlsx")
OUT.mkdir(parents=True,exist_ok=True)

TERMS=[
 "1252961944807641090","1253297784633122817","1429846914028158977",
 "ff2142e98e09b5344994f9beb9c56c95506b9f17",
 "private key","privatekey","weird image","format does not matter",
 "access by private key","mew","puzzle","stegan","hint",
]
def norm(v):
    if v is None: return ""
    if isinstance(v,(datetime,date)): return v.isoformat()
    return str(v)

wb=load_workbook(XLSX,read_only=True,data_only=True)
all_matches=[]
sheet_meta=[]
for ws in wb.worksheets:
    rows=ws.iter_rows(values_only=True)
    try: first=next(rows)
    except StopIteration: continue
    headers=[norm(x) or f"col_{i}" for i,x in enumerate(first)]
    data_rows=[first]
    data_rows.extend(rows)
    count=0
    for ridx,row in enumerate(data_rows,start=1):
        vals=[norm(v) for v in row]
        joined=" | ".join(vals)
        low=joined.lower()
        hits=[t for t in TERMS if t in low]
        date_hit=bool(re.search(r"2020[-/](?:0?4|0?5)[-/]",joined)) or bool(re.search(r"2021[-/]0?8[-/]",joined))
        if not hits and not date_hit: continue
        rec={headers[i] if i<len(headers) else f"col_{i}": vals[i] for i in range(len(vals)) if vals[i]}
        all_matches.append({"sheet":ws.title,"row":ridx,"terms":hits,"date_window":date_hit,"record":rec})
        count+=1
    sheet_meta.append({"sheet":ws.title,"rows":ws.max_row,"cols":ws.max_column,"matches":count})

seen=set(); matches=[]
for m in all_matches:
    key=json.dumps(m["record"],sort_keys=True,ensure_ascii=False)
    if key in seen: continue
    seen.add(key); matches.append(m)

(OUT/"matches.json").write_text(json.dumps({"sheets":sheet_meta,"matches":matches},indent=2,ensure_ascii=False)+"\n")
md=["# Stage 4 — public @ArweaveP archive clue recovery","",
    f"- Workbook sheets scanned: {len(sheet_meta)}",
    f"- Deduplicated matching rows: {len(matches)}","",
    "Source workbook: HomelessPhD/AR_Puzzles/ArweaveP_user_tweets.xlsx.",
    "Rows below are evidence leads from that mirror; they are not primary-source evidence unless a surviving original post can be recovered.",""]
for i,m in enumerate(matches[:250],1):
    md.append(f"## Match {i} — {m['sheet']} row {m['row']}")
    md.append("")
    if m["terms"]: md.append("- Matched: "+", ".join(m["terms"]))
    for k,v in m["record"].items():
        md.append(f"- **{k}**: {v}")
    md.append("")
if len(matches)>250:
    md.append(f"_Report truncated to first 250 matches; JSON contains all {len(matches)}._")
(OUT/"REPORT.md").write_text("\n".join(md)+"\n")
print(json.dumps({"status":"ok","matches":len(matches),"sheets":sheet_meta}))
