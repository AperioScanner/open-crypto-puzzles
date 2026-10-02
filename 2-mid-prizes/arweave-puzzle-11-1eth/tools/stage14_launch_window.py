#!/usr/bin/env python3
"""Stage 14: reconstruct the author's full Puzzle #11 launch-window conversation.

Public archive text mining only; no cryptographic candidate work.
"""
from __future__ import annotations
import json
import re
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"analysis"/"runs"/"stage14-launch-window"
XLSX=Path("/tmp/ArweaveP_user_tweets.xlsx")
OUT.mkdir(parents=True,exist_ok=True)

START=datetime.fromisoformat("2020-04-20T00:00:00")
END=datetime.fromisoformat("2020-05-02T00:00:00")

PATTERNS={
    "321":re.compile(r"(?<!\d)3[\s,./:_-]*2[\s,./:_-]*1(?!\d)",re.I),
    "order":re.compile(r"\b(order|ordered|ordering|sequence|first|second|third|reverse|left|right|clockwise|counterclockwise)\b",re.I),
    "count":re.compile(r"\b(count|counting|number|numbers|how many|three|two|one)\b",re.I),
    "visual":re.compile(r"\b(visual|picture|image|look|overall|environment|format|paper|computer|draw|drawing)\b",re.I),
    "geometry":re.compile(r"\b(perimeter|line|row|column|building|buildings|boat|boats|sail|sails|skyline|harbor|city)\b",re.I),
    "puzzle":re.compile(r"\b(puzzle|stegan|private key|wallet|mew|hidden|hint|solve|solved)\b",re.I),
}

def norm(v):
    if v is None:return ""
    if isinstance(v,datetime):return v.isoformat()
    return str(v)

def parse_dt(s):
    s=(s or "").strip()
    if not s:return None
    for fmt in ("%Y-%m-%dT%H:%M:%S.%fZ","%Y-%m-%dT%H:%M:%SZ","%Y-%m-%dT%H:%M:%S","%Y-%m-%d %H:%M:%S"):
        try:return datetime.strptime(s,fmt)
        except ValueError:pass
    try:return datetime.fromisoformat(s.replace("Z","+00:00")).replace(tzinfo=None)
    except Exception:return None

def main():
    wb=load_workbook(XLSX,read_only=True,data_only=True)
    rows=[]
    for ws in wb.worksheets:
        it=ws.iter_rows(values_only=True)
        try:headers=[norm(x) for x in next(it)]
        except StopIteration:continue
        hm={h.lower():i for i,h in enumerate(headers)}
        ti=hm.get("text"); ui=hm.get("utc")
        if ti is None or ui is None:continue
        for ridx,row in enumerate(it,start=2):
            vals=[norm(v) for v in row]
            text=vals[ti] if ti<len(vals) else ""
            utc=vals[ui] if ui<len(vals) else ""
            dt=parse_dt(utc)
            if dt is None or not (START<=dt<END):continue
            hits=[name for name,pat in PATTERNS.items() if pat.search(text)]
            rec={headers[i] if i<len(headers) else f"col_{i}":vals[i] for i in range(len(vals)) if vals[i]}
            rows.append({"sheet":ws.title,"row":ridx,"utc":dt.isoformat()+"Z","hits":hits,"record":rec})

    rows.sort(key=lambda r:r["utc"])
    # Deduplicate on tweet id/text/time.
    seen=set(); dedup=[]
    for r in rows:
        rec=r["record"]
        key=(rec.get("Tweet Id",""),rec.get("Text",""),r["utc"])
        if key in seen:continue
        seen.add(key); dedup.append(r)

    hit_counts={k:0 for k in PATTERNS}
    for r in dedup:
        for h in r["hits"]:hit_counts[h]+=1

    result={
        "experiment_id":"A11-EXP-014",
        "scope":"all mirrored @ArweaveP posts/replies in 2020-04-20 through 2020-05-01, plus bounded semantic term tagging; public text only",
        "window_start":START.isoformat(),
        "window_end":END.isoformat(),
        "rows":dedup,
        "hit_counts":hit_counts,
        "direct_321_mentions":[r for r in dedup if "321" in r["hits"]],
        "order_mentions":[r for r in dedup if "order" in r["hits"]],
        "count_mentions":[r for r in dedup if "count" in r["hits"]],
        "geometry_mentions":[r for r in dedup if "geometry" in r["hits"]],
    }
    (OUT/"result.json").write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n")

    md=[
        "# Stage 14 — Puzzle #11 launch-window reconstruction","",
        "**Experiment:** A11-EXP-014","",
        "Public archive text mining only. This reconstructs the author's complete mirrored posting/reply window around Puzzle #11 and tags possible ordering/counting/visual clues.","",
        f"- Window: {START.date()} through 2020-05-01 inclusive",
        f"- Deduplicated author posts/replies in window: **{len(dedup)}**",
        f"- Direct 3-2-1 textual patterns: **{hit_counts['321']}**",
        f"- Ordering/direction term hits: **{hit_counts['order']}**",
        f"- Counting/number term hits: **{hit_counts['count']}**",
        f"- Geometry/object term hits: **{hit_counts['geometry']}**","",
        "## Chronological transcript","",
    ]
    for r in dedup:
        rec=r["record"]
        text=rec.get("Text","").replace("\n"," ")
        tid=rec.get("Tweet Id","")
        typ=rec.get("Tweet Type","")
        hits=", ".join(r["hits"]) if r["hits"] else "-"
        md += [f"### {r['utc']} — {tid}",f"- Type: {typ}",f"- Tags: {hits}",f"- Text: {text}",""]
    md += ["## Interpretation",""]
    if hit_counts["321"]:
        md.append("A direct 3-2-1 textual clue exists in the launch window and should be compared against the skyline marker.")
    else:
        md.append("No direct textual 3-2-1 pattern was found in the complete mirrored launch window. This means the skyline 321 marker still lacks an independent textual confirmation from this source.")
    md.append("Ordering/counting/geometry mentions are preserved above for manual contextual review; keyword presence alone is not treated as a clue.")
    md.append("")
    (OUT/"REPORT.md").write_text("\n".join(md))
    print(json.dumps({"status":"ok","experiment_id":"A11-EXP-014","rows":len(dedup),"hit_counts":hit_counts}))

if __name__=="__main__":
    main()
